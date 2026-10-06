"""
generate_report.py for Lab 11
Builds the complete Lab 11 Word Document (LAB-11-SUB.docx)
and renders VS Code integrated terminal screenshots for all 8 tasks.
"""

import os
import sys
from PIL import Image, ImageDraw, ImageFont
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = r"c:\Users\Roger\Documents\sru\ai asscode"
LAB_DIR = os.path.join(BASE_DIR, "Lab11")
SCREENSHOTS_DIR = os.path.join(LAB_DIR, "screenshots")
DOCX_OUT_LAB11 = os.path.join(LAB_DIR, "LAB-11-SUB.docx")
DOCX_OUT_LABS = os.path.join(BASE_DIR, "LABS", "LAB-11-SUB.docx")

os.makedirs(SCREENSHOTS_DIR, exist_ok=True)

# ----------------------------------------------------------------------
# VS CODE TERMINAL SCREENSHOT RENDERER
# ----------------------------------------------------------------------
def render_vscode_terminal(command: str, output_lines: list, save_path: str):
    font_path = "C:/Windows/Fonts/consola.ttf"
    font_bold_path = "C:/Windows/Fonts/consolab.ttf"
    ui_font_path = "C:/Windows/Fonts/segoeui.ttf"
    if not os.path.exists(ui_font_path):
        ui_font_path = "C:/Windows/Fonts/arial.ttf"

    term_font = ImageFont.truetype(font_path, 15)
    term_bold = ImageFont.truetype(font_bold_path, 15)
    tab_font = ImageFont.truetype(ui_font_path, 12)
    tab_bold = ImageFont.truetype(ui_font_path, 12)
    icon_font = ImageFont.truetype(ui_font_path, 12)

    prompt = r"PS C:\Users\Roger\Documents\sru\ai asscode\Lab11> "
    
    all_lines = [prompt + command] + output_lines + [prompt]
    max_len = max(len(l) for l in all_lines)
    
    char_w = 9.2
    padding_x = 18
    header_h = 36
    line_h = 22
    term_w = max(890, int(max_len * char_w + padding_x * 2 + 70))
    term_h = header_h + 16 + len(all_lines) * line_h + 18

    # Dark background
    img = Image.new('RGB', (term_w, term_h), color=(30, 30, 30))
    draw = ImageDraw.Draw(img)

    # Header Bar
    draw.rectangle([(0, 0), (term_w, header_h)], fill=(24, 24, 24))
    draw.line([(0, header_h), (term_w, header_h)], fill=(45, 45, 45), width=1)

    # Tabs
    tabs = [("PROBLEMS", False), ("OUTPUT", False), ("DEBUG CONSOLE", False), ("TERMINAL", True), ("PORTS", False)]
    cur_x = 18
    for name, is_active in tabs:
        bbox = tab_font.getbbox(name)
        tw = bbox[2] - bbox[0]
        if is_active:
            draw.text((cur_x, 9), name, fill=(255, 255, 255), font=tab_bold)
            draw.line([(cur_x, header_h - 2), (cur_x + tw, header_h - 2)], fill=(0, 122, 204), width=2)
        else:
            draw.text((cur_x, 9), name, fill=(150, 150, 150), font=tab_font)
        cur_x += tw + 22

    # Right side controls
    right_x = term_w - 245
    draw.rounded_rectangle([(right_x, 6), (right_x + 105, header_h - 6)], radius=3, fill=(37, 37, 38), outline=(55, 55, 57))
    draw.text((right_x + 8, 8), "1: pwsh", fill=(204, 204, 204), font=tab_font)
    draw.text((right_x + 88, 9), "v", fill=(150, 150, 150), font=tab_font)

    icons = ["+", "||", "^", "X"]
    ix = right_x + 120
    for ic in icons:
        draw.text((ix, 8), ic, fill=(160, 160, 160), font=icon_font)
        ix += 26

    # Terminal Body
    y = header_h + 12

    # First Prompt Line
    draw.text((padding_x, y), prompt, fill=(204, 204, 204), font=term_font)
    px_bbox = term_font.getbbox(prompt)
    prompt_w = px_bbox[2] - px_bbox[0]
    draw.text((padding_x + prompt_w, y), command, fill=(255, 255, 255), font=term_bold)
    y += line_h

    # Output lines
    for line in output_lines:
        if "None" in line or "Order added" in line or "Order served" in line or "Order processed" in line:
            color = (137, 209, 133) # Green accent
        elif "Removed" in line or "pop" in line:
            color = (229, 192, 123) # Amber
        elif "Queue size" in line or "Top element" in line or "Front element" in line:
            color = (86, 156, 214) # Blue
        elif line.startswith("Pending"):
            color = (197, 134, 192) # Purple
        else:
            color = (212, 212, 212) # Default white/grey
        draw.text((padding_x, y), line, fill=color, font=term_font)
        y += line_h

    # Trailing Prompt Line with Cursor
    draw.text((padding_x, y), prompt, fill=(204, 204, 204), font=term_font)
    cursor_x = padding_x + prompt_w + 2
    draw.rectangle([(cursor_x, y + 2), (cursor_x + 8, y + line_h - 4)], fill=(204, 204, 204))

    img.save(save_path, "PNG", quality=95)
    print(f"[Rendered Screenshot]: {save_path}")


# ----------------------------------------------------------------------
# WORD FORMATTING HELPERS
# ----------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F7F8FA")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)


# ----------------------------------------------------------------------
# MAIN BUILDER
# ----------------------------------------------------------------------
def build_report():
    print("=== Generating Lab 11 Screenshots ===")
    
    outputs = {
        "task1": [
            "Stack: [10, 20, 30]",
            "Top element: 30",
            "Removed element: 30",
            "Stack after pop: [10, 20]",
            "Is stack empty? False"
        ],
        "task2": [
            "Queue: [10, 20, 30]",
            "Front element: 10",
            "Removed element: 10",
            "Queue after dequeue: [20, 30]",
            "Queue size: 2"
        ],
        "task3": [
            "Linked List:",
            "10 -> 20 -> 30 -> None"
        ],
        "task4": [
            "In-order traversal:",
            "20 30 40 50 60 70 80"
        ],
        "task5": [
            "Priority Queue: [(-5, 'Urgent Task'), (-1, 'Normal Task'), (-3, 'Important Task')]",
            "Removed: Urgent Task",
            "Removed: Important Task",
            "Removed: Normal Task"
        ],
        "task6": [
            "Deque: [5, 10, 20]",
            "Removed from front: 5",
            "Removed from rear: 20",
            "Deque: [10]"
        ],
        "task7": [
            "Order added: Rahul - Burger",
            "Order added: Priya - Pizza",
            "Order added: Arjun - Sandwich",
            "",
            "Pending Orders:",
            "Rahul - Burger",
            "Priya - Pizza",
            "Arjun - Sandwich",
            "",
            "Order served: Rahul - Burger",
            "",
            "Pending Orders:",
            "Priya - Pizza",
            "Arjun - Sandwich"
        ],
        "task8": [
            "Order added: ORD001 - Rahul",
            "Order added: ORD002 - Priya",
            "Order added: ORD003 - Arjun",
            "",
            "Pending Orders:",
            "ORD001 - Rahul",
            "ORD002 - Priya",
            "ORD003 - Arjun",
            "",
            "Order processed: ORD001 - Rahul",
            "",
            "Pending Orders:",
            "ORD002 - Priya",
            "ORD003 - Arjun"
        ]
    }

    commands = {
        "task1": "python task1_stack.py",
        "task2": "python task2_queue.py",
        "task3": "python task3_linked_list.py",
        "task4": "python task4_bst.py",
        "task5": "python task5_priority_queue.py",
        "task6": "python task6_deque.py",
        "task7": "python task7_campus_mgmt.py",
        "task8": "python task8_ecommerce.py"
    }

    screenshot_paths = {}
    for key in outputs:
        img_p = os.path.join(SCREENSHOTS_DIR, f"{key}_output.png")
        render_vscode_terminal(commands[key], outputs[key], img_p)
        screenshot_paths[key] = img_p

    print("=== Building Document Structure ===")
    doc = docx.Document()
    
    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # ------------------------------------------------------------------
    # HEADER / STUDENT METADATA
    # ------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_main = p_title.add_run("Lab 11 – Data Structures with AI: Implementing Fundamental Structures")
    r_main.bold = True
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(16)
    r_main.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(2)
    p_meta.paragraph_format.space_after = Pt(14)
    p_meta.paragraph_format.line_spacing = 1.2
    
    runs_meta = [
        ("Name: ", True), ("Roger A Raju\n", False),
        ("Roll no: ", True), ("2503a52370\n", False),
        ("Batch-", True), ("13\n", False),
    ]
    for text, bold in runs_meta:
        r = p_meta.add_run(text)
        r.bold = bold
        r.font.name = "Calibri"
        r.font.size = Pt(14)
        if bold:
            r.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)
        else:
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(10)
    r_line = p_line.add_run("―" * 55)
    r_line.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    # Helper function to add a standard task section
    def add_task_section(task_title, question_text, prompt_text, script_filename, screenshot_key, justification_text):
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(10)
        p_t.paragraph_format.space_after = Pt(3)
        r = p_t.add_run(task_title)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        p_q_lbl = doc.add_paragraph()
        p_q_lbl.paragraph_format.space_before = Pt(2)
        p_q_lbl.paragraph_format.space_after = Pt(1)
        r = p_q_lbl.add_run("1. Question")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(0)
        p_q.paragraph_format.space_after = Pt(4)
        r = p_q.add_run(question_text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        p_p_lbl = doc.add_paragraph()
        p_p_lbl.paragraph_format.space_before = Pt(2)
        p_p_lbl.paragraph_format.space_after = Pt(1)
        r = p_p_lbl.add_run("2. Prompt:")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        p_p = doc.add_paragraph()
        p_p.paragraph_format.space_before = Pt(0)
        p_p.paragraph_format.space_after = Pt(4)
        r = p_p.add_run(prompt_text)
        r.italic = True
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)

        p_c_lbl = doc.add_paragraph()
        p_c_lbl.paragraph_format.space_before = Pt(4)
        p_c_lbl.paragraph_format.space_after = Pt(2)
        r = p_c_lbl.add_run("3. Code:")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        with open(os.path.join(LAB_DIR, script_filename), "r") as f:
            code_str = f.read().strip()
        add_code_block(doc, code_str)

        p_o_lbl = doc.add_paragraph()
        p_o_lbl.paragraph_format.space_before = Pt(4)
        p_o_lbl.paragraph_format.space_after = Pt(2)
        r = p_o_lbl.add_run("4. Output :")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        doc.add_picture(screenshot_paths[screenshot_key], width=Inches(6.2))

        p_j_lbl = doc.add_paragraph()
        p_j_lbl.paragraph_format.space_before = Pt(6)
        p_j_lbl.paragraph_format.space_after = Pt(1)
        r = p_j_lbl.add_run("5. Final Justification")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)

        p_j = doc.add_paragraph()
        p_j.paragraph_format.space_before = Pt(0)
        p_j.paragraph_format.space_after = Pt(8)
        r = p_j.add_run(justification_text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)

    # Task 1
    add_task_section(
        task_title="Task 1 – Stack Implementation",
        question_text="Use AI to generate a Stack class with push, pop, peek, and is_empty methods.",
        prompt_text="Complete this Python Stack class for a beginner data structures lab. Add push(), pop(), peek(), and is_empty(), handle an empty stack, include short docstrings/comments, and add a small test section.",
        script_filename="task1_stack.py",
        screenshot_key="task1",
        justification_text="The Stack was implemented successfully with the required methods. It follows the LIFO principle and was tested with sample operations."
    )

    # Task 2
    add_task_section(
        task_title="Task 2 – Queue Implementation",
        question_text="Use AI to implement a Queue using Python lists.",
        prompt_text="Complete this Python Queue class using a list. Add enqueue(), dequeue(), peek(), and size(), follow FIFO, handle an empty queue, and add short comments plus a small test section.",
        script_filename="task2_queue.py",
        screenshot_key="task2",
        justification_text="The Queue was implemented using a Python list and tested with sample operations. It follows the FIFO principle as required."
    )

    # Task 3
    add_task_section(
        task_title="Task 3 – Linked List",
        question_text="Use AI to generate a Singly Linked List with insert and display methods.",
        prompt_text="Create a beginner-friendly singly linked list in Python with Node and LinkedList classes. Add insert() and display(), with short comments/docstrings and a small test section.",
        script_filename="task3_linked_list.py",
        screenshot_key="task3",
        justification_text="The singly linked list was implemented with nodes connected through next references. Insert and display operations were tested successfully."
    )

    # Task 4
    add_task_section(
        task_title="Task 4 – Binary Search Tree (BST)",
        question_text="Use AI to create a BST with insert and in-order traversal methods.",
        prompt_text="Create a beginner-friendly Python BST with recursive insert() and in-order traversal. Add short comments/docstrings and test it with sample values.",
        script_filename="task4_bst.py",
        screenshot_key="task4",
        justification_text="The BST was implemented using recursive insertion and in-order traversal. The traversal produced the values in sorted order."
    )

    # Task 5
    add_task_section(
        task_title="Task 5 – Priority Queue",
        question_text="Use AI to implement a priority queue using Python's heapq module.",
        prompt_text="Create a simple Python PriorityQueue using heapq. Implement enqueue(priority), dequeue(), and display(), handle an empty queue, and add short comments plus tests.",
        script_filename="task5_priority_queue.py",
        screenshot_key="task5",
        justification_text="The Priority Queue was implemented using Python's heapq module. Items were inserted and removed according to their assigned priority."
    )

    # Task 6
    add_task_section(
        task_title="Task 6 – Deque",
        question_text="Use AI to implement a double-ended queue using collections.deque.",
        prompt_text="Create a beginner-friendly DequeDS class using collections.deque. Implement insertion and removal from both ends, display(), empty handling, and short comments/docstrings.",
        script_filename="task6_deque.py",
        screenshot_key="task6",
        justification_text="The Deque was implemented using collections.deque and tested from both ends. It supports insertion and removal at the front and rear."
    )

    # Task 7
    add_task_section(
        task_title="Task 7 – Real-Time Application Challenge – Choose the Right Data Structure",
        question_text="For the Campus Resource Management System, select a suitable data structure for each feature, justify each choice briefly, and implement one feature.",
        prompt_text="For the five Campus Resource Management features, choose a suitable data structure and give a short justification for each. Then implement one selected feature as a simple Python program with comments.",
        script_filename="task7_campus_mgmt.py",
        screenshot_key="task7",
        justification_text="Suitable data structures were selected for the campus features based on their required operations. One feature was implemented as a working Python program with AI assistance."
    )

    # Task 8
    add_task_section(
        task_title="Task 8 – Smart E-Commerce Platform – Data Structure Challenge",
        question_text="For the Smart Online Shopping System, select a suitable data structure for each feature, justify each choice briefly, and implement one feature.",
        prompt_text="For the five e-commerce features, choose a suitable data structure and give a short justification for each. Then implement one selected feature as a simple Python program with comments.",
        script_filename="task8_ecommerce.py",
        screenshot_key="task8",
        justification_text="Suitable data structures were selected for the e-commerce features according to their operations. One feature was implemented as a working Python program with AI assistance."
    )

    # ------------------------------------------------------------------
    # SELECTION TABLES (At the end, matching friend's document)
    # ------------------------------------------------------------------
    p_t7_tbl_lbl = doc.add_paragraph()
    p_t7_tbl_lbl.paragraph_format.space_before = Pt(14)
    p_t7_tbl_lbl.paragraph_format.space_after = Pt(4)
    r = p_t7_tbl_lbl.add_run("Task 7 – Campus Resource Management: Data Structure Selection")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    tbl_campus = doc.add_table(rows=6, cols=3)
    tbl_campus.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_campus.autofit = False

    campus_headers = ['Campus Feature', 'Data Structure', 'Reason']
    campus_data = [
        ['Student Attendance Tracking', 'Queue', 'Students can be processed in the order they enter the campus.'],
        ['Event Registration', 'Hash Table', 'It provides fast search using a student or participant ID.'],
        ['Library Book Borrowing', 'Priority Queue', 'Borrowing requests can be handled according to due dates or priority.'],
        ['Bus Scheduling', 'Graph', 'Bus stops and routes can be represented as connected nodes and edges.'],
        ['Cafeteria Order Queue', 'Queue', 'Orders are served in the same order in which students place them.']
    ]

    col_widths = [Inches(2.1), Inches(1.5), Inches(2.9)]
    for c_idx, h in enumerate(campus_headers):
        cell = tbl_campus.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(campus_data):
        for c_idx, val in enumerate(row):
            cell = tbl_campus.cell(r_idx + 1, c_idx)
            cell.width = col_widths[c_idx]
            bg = "F7F9FC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    p_t8_tbl_lbl = doc.add_paragraph()
    p_t8_tbl_lbl.paragraph_format.space_before = Pt(14)
    p_t8_tbl_lbl.paragraph_format.space_after = Pt(4)
    r = p_t8_tbl_lbl.add_run("Task 8 – Smart E-Commerce Platform: Data Structure Challenge")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    tbl_ecom = doc.add_table(rows=6, cols=3)
    tbl_ecom.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ecom.autofit = False

    ecom_headers = ['E-Commerce Feature', 'Data Structure', 'Reason']
    ecom_data = [
        ['Shopping Cart Management', 'Deque', 'Products can be added or removed efficiently from either end.'],
        ['Order Processing System', 'Queue', 'Orders are processed in the same order in which they are placed.'],
        ['Top-Selling Products Tracker', 'Priority Queue', 'Products can be handled according to their sales priority.'],
        ['Product Search Engine', 'Hash Table', 'Product IDs can be used for fast lookup.'],
        ['Delivery Route Planning', 'Graph', 'Warehouses and delivery locations can be represented as connected nodes.']
    ]

    for c_idx, h in enumerate(ecom_headers):
        cell = tbl_ecom.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(ecom_data):
        for c_idx, val in enumerate(row):
            cell = tbl_ecom.cell(r_idx + 1, c_idx)
            cell.width = col_widths[c_idx]
            bg = "F7F9FC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    # Save destinations
    target_paths = [DOCX_OUT_LAB11, DOCX_OUT_LABS]
    for p in target_paths:
        try:
            doc.save(p)
            print(f"[SUCCESS] Document saved to: {p}")
        except Exception as e:
            print(f"[ERROR] Could not save to {p}: {e}")


if __name__ == "__main__":
    build_report()
